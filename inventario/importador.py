import unicodedata
from decimal import Decimal, InvalidOperation

import openpyxl
from django.db import transaction

from .models import Moneda, Producto, Proveedor

COLUMNAS_ESPERADAS = {
    "codigo": "Código",
    "nombre": "Nombre",
    "descripcion": "Descripción",
    "categoria": "Categoría",
    "unidad_medida": "Unidad de medida",
    "precio_referencial": "Precio referencial",
    "moneda": "Moneda",
    "proveedor": "Proveedor",
}

COLUMNAS_OBLIGATORIAS = {"codigo", "nombre", "precio_referencial"}


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii")
    return texto.strip().lower()


def _mapear_encabezados(fila_encabezados):
    normalizados_esperados = {_normalizar(v): k for k, v in COLUMNAS_ESPERADAS.items()}
    mapa = {}
    for indice, celda in enumerate(fila_encabezados):
        if celda is None:
            continue
        clave = normalizados_esperados.get(_normalizar(celda))
        if clave:
            mapa[clave] = indice
    return mapa


class ResultadoImportacion:
    def __init__(self):
        self.creados = 0
        self.actualizados = 0
        self.proveedores_creados = []
        self.errores = []

    @property
    def total_filas_ok(self):
        return self.creados + self.actualizados


@transaction.atomic
def importar_productos_desde_excel(archivo):
    """`archivo`: ruta o file-like aceptado por openpyxl.load_workbook."""
    resultado = ResultadoImportacion()

    libro = openpyxl.load_workbook(archivo, data_only=True)
    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)

    try:
        encabezados = next(filas)
    except StopIteration:
        resultado.errores.append("El archivo está vacío.")
        return resultado

    mapa = _mapear_encabezados(encabezados)
    faltantes = COLUMNAS_OBLIGATORIAS - mapa.keys()
    if faltantes:
        columnas_legibles = ", ".join(COLUMNAS_ESPERADAS[c] for c in faltantes)
        resultado.errores.append(
            f"Faltan columnas obligatorias en el encabezado: {columnas_legibles}."
        )
        return resultado

    proveedores_cache = {}

    def obtener(fila, clave):
        indice = mapa.get(clave)
        if indice is None or indice >= len(fila):
            return None
        valor = fila[indice]
        if valor is None:
            return None
        valor = str(valor).strip()
        return valor or None

    for numero_fila, fila in enumerate(filas, start=2):
        if fila is None or all(c is None or str(c).strip() == "" for c in fila):
            continue

        codigo = obtener(fila, "codigo")
        nombre = obtener(fila, "nombre")
        if not codigo:
            resultado.errores.append(f"Fila {numero_fila}: falta código.")
            continue
        if not nombre:
            resultado.errores.append(f"Fila {numero_fila}: falta nombre.")
            continue

        precio_texto = obtener(fila, "precio_referencial")
        try:
            precio = Decimal(precio_texto)
            if precio <= 0:
                raise InvalidOperation
        except (InvalidOperation, TypeError):
            resultado.errores.append(
                f"Fila {numero_fila}: precio referencial inválido ('{precio_texto}')."
            )
            continue

        moneda_texto = (obtener(fila, "moneda") or Moneda.PEN).upper()
        if moneda_texto not in Moneda.values:
            resultado.errores.append(
                f"Fila {numero_fila}: moneda no reconocida ('{moneda_texto}'); use PEN o USD."
            )
            continue

        proveedor = None
        nombre_proveedor = obtener(fila, "proveedor")
        if nombre_proveedor:
            clave_proveedor = _normalizar(nombre_proveedor)
            proveedor = proveedores_cache.get(clave_proveedor)
            if proveedor is None:
                proveedor = Proveedor.objects.filter(razon_social__iexact=nombre_proveedor).first()
                if proveedor is None:
                    proveedor = Proveedor.objects.create(razon_social=nombre_proveedor)
                    resultado.proveedores_creados.append(nombre_proveedor)
                proveedores_cache[clave_proveedor] = proveedor

        _, creado = Producto.objects.update_or_create(
            codigo=codigo.upper(),
            defaults={
                "nombre": nombre,
                "descripcion": obtener(fila, "descripcion") or "",
                "categoria": obtener(fila, "categoria") or "",
                "unidad_medida": obtener(fila, "unidad_medida") or "",
                "precio_referencial": precio,
                "moneda": moneda_texto,
                "proveedor_principal": proveedor,
            },
        )
        if creado:
            resultado.creados += 1
        else:
            resultado.actualizados += 1

    return resultado
