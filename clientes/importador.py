import unicodedata

import openpyxl
from django.core.exceptions import ValidationError
from django.core.validators import validate_email
from django.db import transaction

from .forms import validar_ruc
from .models import Cliente

COLUMNAS_ESPERADAS = {
    "codigo_corto": "Código corto",
    "razon_social": "Razón social",
    "ruc": "RUC",
    "direccion": "Dirección",
    "contacto_nombre": "Contacto - nombre",
    "contacto_telefono": "Contacto - teléfono",
    "contacto_email": "Contacto - correo",
}

COLUMNAS_OBLIGATORIAS = {"codigo_corto", "razon_social"}


def _normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto)).encode("ascii", "ignore").decode("ascii")
    return texto.strip().lower()


def _mapear_encabezados(fila_encabezados):
    normalizados_esperados = {_normalizar(v): k for k, v in COLUMNAS_ESPERADAS.items()}
    mapa = {}
    for indice, celda in enumerate(fila_encabezados):
        if celda is None:
            continue
        clave = normalizados_esperados.get(_normalizar(celda))
        if clave:
            mapa[clave] = indice
    return mapa


class ResultadoImportacion:
    def __init__(self):
        self.creados = 0
        self.actualizados = 0
        self.errores = []

    @property
    def total_filas_ok(self):
        return self.creados + self.actualizados


@transaction.atomic
def importar_clientes_desde_excel(archivo):
    """`archivo`: ruta o file-like aceptado por openpyxl.load_workbook.

    Upsert por código corto (no por RUC): es el campo obligatorio y único de verdad en
    Cliente — el RUC es opcional (proveedores/clientes informales pueden no tenerlo), así
    que no sirve como llave confiable para detectar duplicados en todas las filas."""
    resultado = ResultadoImportacion()

    libro = openpyxl.load_workbook(archivo, data_only=True)
    hoja = libro.active
    filas = hoja.iter_rows(values_only=True)

    try:
        encabezados = next(filas)
    except StopIteration:
        resultado.errores.append("El archivo está vacío.")
        return resultado

    mapa = _mapear_encabezados(encabezados)
    faltantes = COLUMNAS_OBLIGATORIAS - mapa.keys()
    if faltantes:
        columnas_legibles = ", ".join(COLUMNAS_ESPERADAS[c] for c in faltantes)
        resultado.errores.append(
            f"Faltan columnas obligatorias en el encabezado: {columnas_legibles}."
        )
        return resultado

    def obtener(fila, clave):
        indice = mapa.get(clave)
        if indice is None or indice >= len(fila):
            return None
        valor = fila[indice]
        if valor is None:
            return None
        valor = str(valor).strip()
        return valor or None

    for numero_fila, fila in enumerate(filas, start=2):
        if fila is None or all(c is None or str(c).strip() == "" for c in fila):
            continue

        codigo_corto = obtener(fila, "codigo_corto")
        razon_social = obtener(fila, "razon_social")
        if not codigo_corto:
            resultado.errores.append(f"Fila {numero_fila}: falta código corto.")
            continue
        if not razon_social:
            resultado.errores.append(f"Fila {numero_fila}: falta razón social.")
            continue

        ruc = obtener(fila, "ruc") or ""
        if ruc:
            try:
                validar_ruc(ruc)
            except ValidationError:
                resultado.errores.append(f"Fila {numero_fila}: RUC inválido ('{ruc}').")
                continue

        contacto_email = obtener(fila, "contacto_email") or ""
        if contacto_email:
            try:
                validate_email(contacto_email)
            except ValidationError:
                resultado.errores.append(f"Fila {numero_fila}: correo inválido ('{contacto_email}').")
                continue

        _, creado = Cliente.objects.update_or_create(
            codigo_corto=codigo_corto.upper(),
            defaults={
                "razon_social": razon_social,
                "ruc": ruc,
                "direccion": obtener(fila, "direccion") or "",
                "contacto_nombre": obtener(fila, "contacto_nombre") or "",
                "contacto_telefono": obtener(fila, "contacto_telefono") or "",
                "contacto_email": contacto_email,
            },
        )
        if creado:
            resultado.creados += 1
        else:
            resultado.actualizados += 1

    return resultado
